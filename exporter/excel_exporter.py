from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

_COLUMNS = [
    ("_row_num",           "#",                   5),
    ("name",               "Business Name",       32),
    ("address",            "Address",             38),
    ("phone",              "Phone",               16),
    ("website",            "Website URL",         42),
    ("website_status",     "Website Status",      16),
    ("last_evidence_date", "Last Evidence Date",  22),
    ("method_used",        "Detection Method",    22),
    ("rating",             "Rating",               9),
    ("postcode",           "Postcode Searched",   16),
]

_STATUS_FILLS = {
    "No Website": PatternFill(fill_type="solid", fgColor="FFD7D7"),
    "Outdated":   PatternFill(fill_type="solid", fgColor="FFE5B4"),
    "OK":         PatternFill(fill_type="solid", fgColor="D7FFD7"),
    "Unknown":    PatternFill(fill_type="solid", fgColor="E8E8E8"),
}

_THIN = Side(border_style="thin", color="BBBBBB")
_MEDIUM = Side(border_style="medium", color="888888")
_THIN_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_MEDIUM)


class ExcelExporter:
    def __init__(self, results: list[dict]) -> None:
        self._results = results

    def export(self, filepath: str) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Results"

        self._write_metadata(ws)
        self._write_headers(ws, row=4)
        self._write_data(ws, start_row=5)
        self._apply_column_widths(ws)

        last_col = get_column_letter(len(_COLUMNS))
        ws.freeze_panes = "A5"
        ws.auto_filter.ref = f"A4:{last_col}4"

        wb.save(filepath)

    def _write_metadata(self, ws) -> None:
        last_col = get_column_letter(len(_COLUMNS))

        ws.merge_cells(f"A1:{last_col}1")
        cell = ws["A1"]
        cell.value = "Google Maps Business Scraper — Results"
        cell.font = Font(bold=True, size=14, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

        ws.merge_cells(f"A2:{last_col}2")
        cell = ws["A2"]
        cell.value = (
            f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')} "
            f"| Results: {len(self._results)}"
        )
        cell.font = Font(italic=True, size=10, color="1F4E79")
        cell.fill = PatternFill(fill_type="solid", fgColor="D6E4F0")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20

        ws.row_dimensions[3].height = 6

    def _write_headers(self, ws, row: int) -> None:
        header_fill = PatternFill(fill_type="solid", fgColor="2E75B6")
        header_font = Font(bold=True, size=11, color="FFFFFF")
        centred = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_idx, (_, label, _) in enumerate(_COLUMNS, start=1):
            cell = ws.cell(row=row, column=col_idx, value=label)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = centred
            cell.border = _HEADER_BORDER
        ws.row_dimensions[row].height = 30

    def _write_data(self, ws, start_row: int) -> None:
        url_font = Font(color="0563C1", underline="single", size=10)
        default_font = Font(size=10)
        default_align = Alignment(vertical="center", wrap_text=False)
        address_align = Alignment(vertical="center", wrap_text=True)
        centre_align = Alignment(horizontal="center", vertical="center")

        for row_num, result in enumerate(self._results, start=1):
            row = start_row + row_num - 1
            status = result.get("website_status", "Unknown")
            row_fill = _STATUS_FILLS.get(status, _STATUS_FILLS["Unknown"])

            values = [
                row_num,
                result.get("name", ""),
                result.get("address", ""),
                result.get("phone", ""),
                result.get("website", ""),
                status,
                result.get("last_evidence_date", ""),
                result.get("method_used", ""),
                result.get("rating", ""),
                result.get("postcode", ""),
            ]

            for col_idx, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col_idx, value=value)
                cell.fill = row_fill
                cell.border = _THIN_BORDER

                col_key = _COLUMNS[col_idx - 1][0]

                if col_key == "website" and value:
                    cell.font = url_font
                    cell.hyperlink = value
                    cell.alignment = default_align
                elif col_key in ("_row_num", "website_status", "rating", "postcode"):
                    cell.font = default_font
                    cell.alignment = centre_align
                elif col_key == "address":
                    cell.font = default_font
                    cell.alignment = address_align
                else:
                    cell.font = default_font
                    cell.alignment = default_align

            ws.row_dimensions[row].height = 22

    def _apply_column_widths(self, ws) -> None:
        for col_idx, (_, _, width) in enumerate(_COLUMNS, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width
